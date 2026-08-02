from io import BytesIO

from openpyxl import load_workbook

from app.reports import report_payload
from app.schemas import PointTrack


TRACKS = [PointTrack(frame=1, timestampMs=33, x=10.25, y=20.5, model="circle", residual=.03, confidence=.94, state="valid", durationMs=1.2), PointTrack(frame=0, timestampMs=0, x=10.1, y=20.3, model="circle", residual=.04, confidence=.92, state="valid", durationMs=1.4)]


def test_xlsx_contains_tracks_and_events() -> None:
    content, _, _ = report_payload("job-report", TRACKS, "xlsx")
    workbook = load_workbook(BytesIO(content), read_only=True)
    assert workbook.sheetnames == ["tracks", "events"]
    assert workbook["tracks"]["A2"].value == 0


def test_empty_report_is_recoverable_error() -> None:
    try:
        report_payload("empty", [], "pdf")
    except ValueError as error:
        assert str(error) == "export.empty"
    else:
        raise AssertionError("expected export.empty")
